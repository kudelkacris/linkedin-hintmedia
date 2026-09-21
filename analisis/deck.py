# -*- coding: utf-8 -*-
"""Helper de estilo para los decks Hint Media."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BG      = RGBColor(0x0E, 0x11, 0x16)
CARD    = RGBColor(0x18, 0x1D, 0x25)
CARD2   = RGBColor(0x22, 0x28, 0x33)
TXT     = RGBColor(0xF2, 0xF5, 0xF8)
MUT     = RGBColor(0x8A, 0x94, 0xA6)
ACC     = RGBColor(0x5B, 0x9D, 0xF9)   # azul
GOOD    = RGBColor(0x3D, 0xD6, 0x8C)   # verde
BAD     = RGBColor(0xF8, 0x71, 0x71)   # rojo
WARN    = RGBColor(0xFB, 0xBF, 0x24)   # amarillo
GOLD    = RGBColor(0xE8, 0xB4, 0x5B)

W, H = Inches(13.333), Inches(7.5)


def new_deck():
    p = Presentation()
    p.slide_width, p.slide_height = W, H
    return p


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bgf = s.background.fill
    bgf.solid()
    bgf.fore_color.rgb = BG
    return s


def tb(slide, x, y, w, h, text, size=18, color=TXT, bold=False, align=PP_ALIGN.LEFT,
       font="Segoe UI", anchor=MSO_ANCHOR.TOP, space=0, italic=False, line=None):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        if line:
            p.line_spacing = line
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
    return box


def rect(slide, x, y, w, h, fill=CARD, lineclr=None, radius=True, lw=1.2):
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, x, y, w, h)
    if radius:
        try:
            sh.adjustments[0] = 0.06
        except Exception:
            pass
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if lineclr is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = lineclr
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def title(slide, kicker, head, sub=None):
    tb(slide, Inches(0.7), Inches(0.42), Inches(11), Inches(0.3), kicker.upper(),
       size=11, color=ACC, bold=True)
    tb(slide, Inches(0.7), Inches(0.72), Inches(12), Inches(0.62), head, size=30, bold=True)
    if sub:
        tb(slide, Inches(0.7), Inches(1.38), Inches(12), Inches(0.4), sub, size=13, color=MUT)
    ln = rect(slide, Inches(0.7), Inches(1.28) if not sub else Inches(1.84),
              Inches(1.1), Pt(3), fill=ACC, radius=False)
    return ln


def kpi(slide, x, y, w, big, label, sub=None, color=TXT, h=Inches(1.5)):
    rect(slide, x, y, w, h, fill=CARD)
    tb(slide, x + Inches(0.22), y + Inches(0.16), w - Inches(0.44), Inches(0.62), big,
       size=34, bold=True, color=color)
    tb(slide, x + Inches(0.22), y + Inches(0.84), w - Inches(0.44), Inches(0.3), label,
       size=11.5, color=TXT, bold=True)
    if sub:
        tb(slide, x + Inches(0.22), y + Inches(1.1), w - Inches(0.44), Inches(0.3), sub,
           size=9.5, color=MUT)


def bars(slide, x, y, w, h, data, maxv=None, fmt="%.1f%%", barh=Inches(0.3),
         gap=Inches(0.17), labw=Inches(2.2), colorfn=None, valw=Inches(0.95)):
    """data = [(label, value, color|None)]"""
    mx = maxv or max([d[1] for d in data] + [0.0001])
    cy = y
    for item in data:
        lab, val = item[0], item[1]
        clr = item[2] if len(item) > 2 and item[2] else (colorfn(val) if colorfn else ACC)
        tb(slide, x, cy - Inches(0.02), labw, barh, lab, size=11, color=TXT,
           anchor=MSO_ANCHOR.MIDDLE)
        track_w = w - labw - valw
        rect(slide, x + labw, cy, track_w, barh, fill=CARD2)
        bw = int(track_w * (val / mx)) if mx else 0
        if bw > 0:
            rect(slide, x + labw, cy, max(bw, Emu(6000)), barh, fill=clr)
        tb(slide, x + labw + track_w + Inches(0.12), cy - Inches(0.02), valw, barh,
           fmt % val, size=11.5, color=clr, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        cy += barh + gap
    return cy


def table(slide, x, y, w, headers, rowsdata, colw=None, size=11, rowh=Inches(0.34),
          headclr=MUT, zebra=True, colorcols=None):
    n = len(headers)
    colw = colw or [w / n] * n
    cy = y
    cx = x
    for i, hd in enumerate(headers):
        tb(slide, cx, cy, colw[i], Inches(0.3), hd, size=10, color=headclr, bold=True,
           align=PP_ALIGN.RIGHT if i else PP_ALIGN.LEFT)
        cx += colw[i]
    cy += Inches(0.32)
    rect(slide, x, cy, w, Pt(1), fill=CARD2, radius=False)
    cy += Inches(0.1)
    for ri, row in enumerate(rowsdata):
        if zebra and ri % 2 == 0:
            rect(slide, x - Inches(0.1), cy - Inches(0.04), w + Inches(0.2), rowh, fill=CARD)
        cx = x
        for i, cell in enumerate(row):
            clr = TXT
            bold = (i == 0)
            if colorcols and i in colorcols:
                clr = colorcols[i](cell, ri)
                bold = True
            tb(slide, cx, cy, colw[i], rowh, str(cell), size=size, color=clr, bold=bold,
               align=PP_ALIGN.RIGHT if i else PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
            cx += colw[i]
        cy += rowh
    return cy


def bullets(slide, x, y, w, items, size=13, gap=Inches(0.52), dotclr=ACC, sub_size=11):
    cy = y
    for it in items:
        if isinstance(it, tuple):
            head, sub = it
        else:
            head, sub = it, None
        rect(slide, x, cy + Inches(0.09), Inches(0.1), Inches(0.1), fill=dotclr, radius=True)
        tb(slide, x + Inches(0.28), cy, w - Inches(0.28), Inches(0.3), head, size=size, bold=True)
        if sub:
            tb(slide, x + Inches(0.28), cy + Inches(0.27), w - Inches(0.28), Inches(0.4), sub,
               size=sub_size, color=MUT, line=1.25)
            cy += gap + Inches(0.16)
        else:
            cy += gap - Inches(0.14)
    return cy


def footer(slide, txt, num=None):
    tb(slide, Inches(0.7), Inches(6.95), Inches(9), Inches(0.3), txt, size=9, color=MUT)
    if num is not None:
        tb(slide, Inches(12.0), Inches(6.95), Inches(0.63), Inches(0.3), str(num), size=9,
           color=MUT, align=PP_ALIGN.RIGHT)
